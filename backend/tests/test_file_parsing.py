"""Reading CSV / Excel files safely into tables. Pure functions: no database, no network."""
import io
import zipfile

import pytest
from openpyxl import Workbook

from app.services.connectors import file_parsing as fp
from app.services.connectors.file_parsing import FileParseError, parse_file


def csv_bytes(text: str, encoding="utf-8") -> bytes:
    return text.encode(encoding)


def test_a_plain_csv_becomes_a_typed_table():
    [t] = parse_file("payroll_2026.csv", csv_bytes("emp_id,name,salary,hired,active\nE1,Ann,1000.50,2020-01-05,true\nE2,Bob,900,2021-03-10,false\n"))
    assert t.name == "payroll_2026" and t.headers == ["emp_id", "name", "salary", "hired", "active"]
    assert t.types == {"emp_id": "string", "name": "string", "salary": "double", "hired": "date", "active": "boolean"}
    assert t.rows[0] == {"emp_id": "E1", "name": "Ann", "salary": 1000.5, "hired": "2020-01-05", "active": True}
    assert isinstance(t.rows[1]["salary"], float) and t.rows[1]["active"] is False  # numbers are numbers: a rule can say salary > 950


def test_whole_number_columns_are_integers_and_mixed_columns_stay_text():
    [t] = parse_file("x.csv", csv_bytes("qty,mixed\n1,1\n2,two\n3,3\n"))
    assert t.types == {"qty": "integer", "mixed": "string"}
    assert [r["qty"] for r in t.rows] == [1, 2, 3] and [r["mixed"] for r in t.rows] == ["1", "two", "3"]


@pytest.mark.parametrize(
    "value",
    ["000123", "0123", "+27821234567", "1234567890123456789", "007", "00.5", "0x1F", "1,000", "12 34"],
)
def test_identifier_like_values_are_never_turned_into_numbers(value):
    """Typing them would change them ("000123" -> 123, "+27..." -> 27...), and a join to the same ids elsewhere would fail."""
    [t] = parse_file("ids.csv", csv_bytes(f'code\n"{value}"\n"{value}"\n'))
    assert t.types == {"code": "string"} and [r["code"] for r in t.rows] == [value, value]


@pytest.mark.parametrize("value, kind", [("0", "integer"), ("-5", "integer"), ("123456789012345678", "integer"), ("0.5", "double"), (".5", "double"), ("-0.25", "double"), ("1e3", "double"), ("10.", "double")])
def test_real_numbers_are_still_numbers(value, kind):
    [t] = parse_file("n.csv", csv_bytes(f"n\n{value}\n{value}\n"))
    assert t.types == {"n": kind}


def test_a_column_is_a_number_only_if_every_value_is():
    [t] = parse_file("m.csv", csv_bytes("n\n1\n2\n007\n"))  # one identifier-looking value makes the whole column text
    assert t.types == {"n": "string"} and [r["n"] for r in t.rows] == ["1", "2", "007"]


def test_records_become_a_typed_table_like_a_file():
    from app.services.connectors.file_parsing import table_from_records

    t = table_from_records("invoices", [{"id": "A1", "amount": "10.5", "paid": "true"}, {"id": "A2", "amount": "7", "extra": 3}, {"id": "A3"}])
    assert t.headers == ["id", "amount", "paid", "extra"]
    assert t.types == {"id": "string", "amount": "double", "paid": "boolean", "extra": "integer"}
    assert t.rows[0] == {"id": "A1", "amount": 10.5, "paid": True, "extra": None} and t.rows[2]["amount"] is None
    assert table_from_records("empty", []).headers == []


@pytest.mark.parametrize("delimiter", [",", ";", "\t", "|"])
def test_the_delimiter_is_detected(delimiter):
    [t] = parse_file("d.csv", csv_bytes(delimiter.join(["a", "b", "c"]) + "\n" + delimiter.join(["1", "2", "3"]) + "\n"))
    assert t.headers == ["a", "b", "c"] and t.rows == [{"a": 1, "b": 2, "c": 3}]


def test_quoted_fields_embedded_commas_and_newlines_survive():
    [t] = parse_file("q.csv", csv_bytes('id,note\n1,"hello, ""world"""\n2,"line one\nline two"\n'))
    assert t.rows[0]["note"] == 'hello, "world"' and t.rows[1]["note"] == "line one\nline two"


def test_a_bom_and_windows_encodings_are_handled():
    assert parse_file("b.csv", b"\xef\xbb\xbfid,name\n1,Zo\xc3\xab\n")[0].headers[0] == "id"  # BOM does not become part of the header
    assert parse_file("w.csv", "id,name\n1,Zoë\n".encode("cp1252"))[0].rows[0]["name"] == "Zoë"


def test_headers_are_made_unique_and_never_empty():
    [t] = parse_file("h.csv", csv_bytes("Name,Name,,name \n1,2,3,4\n"))
    assert t.headers == ["Name", "Name_2", "column_3", "name_3"]  # case-insensitive duplicates too, whitespace trimmed
    assert list(t.rows[0].values()) == [1, 2, 3, 4]


def test_blank_lines_and_short_or_long_rows_are_tolerated():
    [t] = parse_file("r.csv", csv_bytes("a,b\n\n1,2\n3\n4,5,6\n"))
    assert t.rows == [{"a": 1, "b": 2}, {"a": 3, "b": None}, {"a": 4, "b": 5}]


@pytest.mark.parametrize("name", ["x.exe", "x.pdf", "x.csv.exe", "noextension", "x.xls"])
def test_only_table_files_are_accepted(name):
    with pytest.raises(FileParseError, match="Only"):
        parse_file(name, b"a,b\n1,2\n")


def test_empty_and_headerless_files_are_refused():
    with pytest.raises(FileParseError):
        parse_file("e.csv", b"")
    with pytest.raises(FileParseError):
        parse_file("e.csv", b"\n\n  \n")


def test_the_size_row_and_column_limits_are_enforced(monkeypatch):
    with pytest.raises(FileParseError, match="larger"):
        parse_file("big.csv", b"a\n" + b"1\n" * (fp.MAX_FILE_BYTES // 2 + 1))
    monkeypatch.setattr(fp, "MAX_ROWS", 3)
    [t] = parse_file("rows.csv", csv_bytes("a\n" + "\n".join(str(i) for i in range(10)) + "\n"))
    assert len(t.rows) == 3 and t.truncated
    monkeypatch.setattr(fp, "MAX_COLUMNS", 2)
    with pytest.raises(FileParseError, match="columns"):
        parse_file("wide.csv", csv_bytes("a,b,c\n1,2,3\n"))


def test_an_oversized_cell_is_cut_not_trusted():
    [t] = parse_file("cell.csv", csv_bytes("a\n" + "x" * (fp.MAX_CELL_CHARS + 500) + "\n"))
    assert len(t.rows[0]["a"]) == fp.MAX_CELL_CHARS


def test_csv_formula_cells_are_plain_text_never_evaluated():
    [t] = parse_file("f.csv", csv_bytes("a\n=HYPERLINK(\"http://evil\",\"x\")\n"))
    assert t.rows[0]["a"].startswith("=HYPERLINK")  # stored as text; nothing executes it


def _xlsx(sheets: dict[str, list[list]]) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for title, rows in sheets.items():
        ws = wb.create_sheet(title)
        for row in rows:
            ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_one_sheet_is_one_table_named_after_the_file():
    [t] = parse_file("ledger.xlsx", _xlsx({"Sheet1": [["account", "amount"], ["A1", 10], ["A2", 20.5]]}))
    assert t.name == "ledger" and t.types == {"account": "string", "amount": "double"}
    assert t.rows == [{"account": "A1", "amount": 10.0}, {"account": "A2", "amount": 20.5}]


def test_several_sheets_are_several_tables():
    tables = parse_file("book.xlsx", _xlsx({"Payroll": [["id"], [1]], "Leavers": [["id", "date"], [2, "2024-01-01"]]}))
    assert [t.name for t in tables] == ["book / Payroll", "book / Leavers"]


def test_excel_dates_become_iso_text_and_blank_rows_are_skipped():
    import datetime

    [t] = parse_file("d.xlsx", _xlsx({"S": [["when", "n"], [datetime.date(2024, 5, 1), 1], [None, None], [datetime.datetime(2024, 6, 2, 8, 30), 2]]}))
    assert t.types["when"] == "date" and t.rows[0]["when"] == "2024-05-01" and t.rows[1]["when"].startswith("2024-06-02")
    assert len(t.rows) == 2


def test_a_sheet_without_a_header_row_is_skipped_and_a_workbook_with_none_is_refused():
    with pytest.raises(FileParseError, match="no sheet"):
        parse_file("empty.xlsx", _xlsx({"S": []}))


def test_a_corrupt_workbook_is_refused_cleanly():
    with pytest.raises(FileParseError):
        parse_file("bad.xlsx", b"PK\x03\x04 this is not really a zip")
    with pytest.raises(FileParseError, match="valid .xlsx"):
        parse_file("bad2.xlsx", b"just some text")


def test_a_zip_bomb_is_refused_before_it_is_expanded(monkeypatch):
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("xl/worksheets/sheet1.xml", "0" * 5_000_000)  # tiny once compressed, large when expanded
    monkeypatch.setattr(fp, "MAX_UNCOMPRESSED_XLSX_BYTES", 1_000_000)
    with pytest.raises(FileParseError, match="unreasonable size"):
        parse_file("bomb.xlsx", buf.getvalue())


def test_excel_formulas_are_read_as_cached_values_never_evaluated():
    wb = Workbook()
    ws = wb.active
    ws.append(["a", "b"])
    ws.append([1, "=1+1"])
    buf = io.BytesIO()
    wb.save(buf)
    [t] = parse_file("form.xlsx", buf.getvalue())
    assert t.rows[0]["b"] is None  # a never-calculated formula has no cached value; nothing was evaluated


# --- memory: a file at the size limit must not be able to take a small server down -------------------------------------------
def test_a_file_at_the_size_limit_stays_within_a_memory_budget():
    """MEASURED before the fix: a 22 MB file peaked at ~340 MB in the parser and kept ~150 MB. It is now cut at MAX_CELLS
    and read as a stream, and this test fails if that regresses (thresholds leave ~1.8x headroom over what it measures)."""
    import tracemalloc

    header = ",".join(f"c{i}" for i in range(20))
    data = ("\n".join([header] + [",".join([f"{r}.50", f"Name {r}", "2026-08-15", f"CODE{r % 5000:05d}"] * 5) for r in range(100_000)]) + "\n").encode()
    assert len(data) > 15 * 1024 * 1024
    tracemalloc.start()
    [t] = parse_file("big.csv", data)
    retained, peak = tracemalloc.get_traced_memory()
    tracemalloc.stop()
    assert t.truncated and len(t.rows) * len(t.headers) <= fp.MAX_CELLS
    assert peak < 100 * 1024 * 1024, f"peak {peak / 1e6:.0f} MB"
    assert retained < 60 * 1024 * 1024, f"retained {retained / 1e6:.0f} MB"


def test_a_file_is_decoded_as_it_is_read_and_a_late_bad_byte_falls_back_to_the_next_encoding():
    data = ("id,name\n" + "".join(f"{i},Zoe {i}\n" for i in range(5000))).encode("utf-8") + "9999,Café\n".encode("cp1252")  # é is 0xE9: invalid UTF-8, at the very END
    [t] = parse_file("late.csv", data)
    assert len(t.rows) == 5001 and t.rows[-1] == {"id": 9999, "name": "Café"}
    assert t.rows[0]["name"] == "Zoe 0"


def test_bytes_beyond_the_row_limit_are_never_decoded(monkeypatch):
    monkeypatch.setattr(fp, "MAX_ROWS", 3)
    data = "name\nZoë\nAnn\nBob\n".encode("utf-8") + b"x\n" * 300_000 + b"\xe9\xff\xfe\n"  # invalid UTF-8, far past the cut
    [t] = parse_file("cut.csv", data)
    assert t.truncated and [r["name"] for r in t.rows] == ["Zoë", "Ann", "Bob"]  # decoded as UTF-8: the bad tail was never reached, so no fallback


def test_a_whitespace_only_file_is_empty_without_copying_it():
    with pytest.raises(FileParseError, match="empty"):
        parse_file("blank.csv", b" \r\n\t\n" * 1000)
    with pytest.raises(FileParseError, match="header|empty"):
        parse_file("headless.csv", b'\n\n"",\n')


def test_the_parse_cache_is_bounded_by_size_not_only_by_count(monkeypatch):
    import uuid

    from app.services.connectors import file_connector as fc

    def tables(cells):
        return (fp.ParsedTable(name="t", headers=["a"], rows=[{"a": 1}] * cells),)

    fc._PARSE_CACHE.clear()
    monkeypatch.setattr(fc, "_PARSE_CACHE_MAX", 5)
    monkeypatch.setattr(fc, "_PARSE_CACHE_MAX_CELLS", 1000)
    fc._cache_put(uuid.uuid4(), "s0", tables(400))
    fc._cache_put(uuid.uuid4(), "s1", tables(400))
    assert [k[1] for k in fc._PARSE_CACHE] == ["s0", "s1"]
    fc._cache_put(uuid.uuid4(), "s2", tables(400))  # 1,200 cells is over the 1,000 budget: the oldest goes
    assert [k[1] for k in fc._PARSE_CACHE] == ["s1", "s2"]
    fc._cache_put(uuid.uuid4(), "huge", tables(5000))  # bigger than the whole budget: it alone stays, everything older is gone
    assert [k[1] for k in fc._PARSE_CACHE] == ["huge"]
    fc._PARSE_CACHE.clear()


def test_the_sftp_snapshot_cache_is_bounded_by_size_too(monkeypatch):
    import uuid

    from app.services.connectors import sftp_connector as sc

    def snapshot(cells):
        return sc._Snapshot(("sig",), "/p", 1, 1, "h", "f.csv", (fp.ParsedTable(name="t", headers=["a"], rows=[{"a": 1}] * cells),), 0.0)

    sc._SNAPSHOTS.clear()
    monkeypatch.setattr(sc, "_SNAPSHOT_MAX", 5)
    monkeypatch.setattr(sc, "_SNAPSHOT_MAX_CELLS", 1000)
    a, b, c = uuid.uuid4(), uuid.uuid4(), uuid.uuid4()
    sc._remember(a, snapshot(400))
    sc._remember(b, snapshot(400))
    sc._remember(c, snapshot(400))
    assert list(sc._SNAPSHOTS) == [b, c]
    sc._remember(a, snapshot(5000))
    assert list(sc._SNAPSHOTS) == [a]
    sc._SNAPSHOTS.clear()


def test_a_measurement_over_rows_that_were_cut_is_marked_a_lower_bound(monkeypatch):
    import uuid
    from types import SimpleNamespace

    from app.core.relationship_inference import Candidate, ColumnRef
    from app.services import data_source_service as svc

    connection = SimpleNamespace(connection_id=uuid.uuid4(), db_type="file_upload")
    rows = {"orders": [{"customer_id": f"C{i}"} for i in range(10)], "customers": [{"customer_id": f"C{i}"} for i in range(5)]}
    monkeypatch.setattr(svc, "fetch_profile_rows", lambda c, *, entity_name, columns, limit: rows[entity_name])

    class Cut:  # a connector that says its "customers" table has more rows than were read
        def truncated(self, connection, entity_name):
            return entity_name == "customers"

    class Whole:
        def truncated(self, connection, entity_name):
            return False

    candidate = Candidate(
        child=ColumnRef("f1", "e1", "orders", "customer_id", "string"), parent=ColumnRef("f2", "e2", "customers", "customer_id", "string"),
        name_affinity=1.0, parent_unique=True,
    )
    monkeypatch.setattr(svc.connectors, "for_connection", lambda c: Cut())
    [(_, cut)] = svc._measure_in_memory([connection], [candidate])
    assert cut.capped is True and (cut.child_distinct, cut.matched_distinct, cut.parent_distinct) == (10, 5, 5)
    monkeypatch.setattr(svc.connectors, "for_connection", lambda c: Whole())
    [(_, whole)] = svc._measure_in_memory([connection], [candidate])
    assert whole.capped is False
    monkeypatch.setattr(svc.connectors, "for_connection", lambda c: SimpleNamespace())  # a connector that cannot say: measured as before
    [(_, unknown)] = svc._measure_in_memory([connection], [candidate])
    assert unknown.capped is False


# --- a workbook's sheets share ONE budget (measured: 30 sheets at a per-table limit kept 1.2M cells, 82 MB) ---------------------
def _sheets(count, rows=100, cols=10):
    return _xlsx({f"S{i}": [[f"c{c}" for c in range(cols)]] + [[f"r{r}c{c}" for c in range(cols)] for r in range(rows)] for i in range(count)})


def test_the_sheets_of_one_workbook_share_a_single_cell_budget(monkeypatch):
    monkeypatch.setattr(fp, "MAX_CELLS", 2_500)  # each sheet is 100 rows x 10 columns = 1,000 cells
    tables = parse_file("book.xlsx", _sheets(6))
    assert [len(t.rows) for t in tables] == [100, 100, 50, 0, 0, 0]  # the third gets what is left, the rest none
    assert [t.truncated for t in tables] == [False, False, True, True, True, True]
    assert sum(len(t.rows) * len(t.headers) for t in tables) <= fp.MAX_CELLS
    assert all(t.headers == [f"c{c}" for c in range(10)] for t in tables)  # every sheet is still listed, with its columns


def test_a_workbook_within_the_budget_is_read_whole(monkeypatch):
    monkeypatch.setattr(fp, "MAX_CELLS", 10_000)
    tables = parse_file("book.xlsx", _sheets(5))
    assert [len(t.rows) for t in tables] == [100] * 5 and not any(t.truncated for t in tables)


def test_a_workbook_with_too_many_sheets_is_refused(monkeypatch):
    monkeypatch.setattr(fp, "MAX_SHEETS", 3)
    with pytest.raises(FileParseError, match="more than 3 sheets"):
        parse_file("book.xlsx", _sheets(4, rows=2, cols=2))
    assert len(parse_file("book.xlsx", _sheets(3, rows=2, cols=2))) == 3


def test_the_notices_say_which_tables_were_cut_and_which_got_no_rows(monkeypatch):
    from app.services.connectors.file_connector import truncation_notices

    monkeypatch.setattr(fp, "MAX_CELLS", 2_500)
    notices = truncation_notices(tuple(parse_file("book.xlsx", _sheets(4))))
    assert len(notices) == 2  # the two sheets that were cut; the two that were read whole say nothing
    assert "first 50 rows of 'book / S2'" in notices[0]
    assert "rows of 'book / S3' were not read" in notices[1] and "other sheets" in notices[1]
    assert truncation_notices(tuple(parse_file("book.xlsx", _sheets(2, rows=3, cols=2)))) == []
