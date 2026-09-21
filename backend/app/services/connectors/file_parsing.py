"""
Reading a CSV / TSV / Excel file into tables the platform can discover, map and test.

Pure functions over bytes (no database, no network), so uploads and SFTP fetches share exactly one
parser and it can be tested exhaustively. Everything that reads a stranger's file is bounded: size,
rows, columns, cell length, and (for Excel, which is a zip) the decompressed size.

A "table" is one CSV, or one Excel sheet. Values are coerced per COLUMN to the type the whole column
supports (all whole numbers -> integer, all numeric -> double, all ISO dates -> date, all true/false ->
boolean, otherwise text), so a rule comparing `amount > 1000` compares numbers, not strings. The type
labels match what the MongoDB connector reports, which the value-profile code already understands.
"""
from __future__ import annotations

import csv
import io
import re
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime

MAX_FILE_BYTES = 25 * 1024 * 1024
MAX_ROWS = 100_000
MAX_COLUMNS = 500
# Rows x columns is what actually costs memory (each cell is a Python object): 100k rows of 500 columns
# would be ~50M of them. A wide file is cut to fewer rows instead, and reported as truncated.
MAX_CELLS = 2_000_000
MAX_CELL_CHARS = 32_000
MAX_UNCOMPRESSED_XLSX_BYTES = 200 * 1024 * 1024  # a zip bomb inflates far past its file size

_CSV_SUFFIXES = (".csv", ".tsv", ".txt")
_XLSX_SUFFIXES = (".xlsx", ".xlsm")


class FileParseError(ValueError):
    """The file cannot be read as a table. The message is safe to show the user."""


@dataclass
class ParsedTable:
    name: str
    headers: list[str]
    rows: list[dict]
    types: dict[str, str] = field(default_factory=dict)  # header -> integer|double|date|boolean|string
    truncated: bool = False


def file_kind(file_name: str) -> str:
    lowered = file_name.lower()
    if lowered.endswith(_XLSX_SUFFIXES):
        return "excel"
    if lowered.endswith(_CSV_SUFFIXES):
        return "csv"
    raise FileParseError("Only .csv, .tsv, .txt and .xlsx files are supported.")


def table_stem(file_name: str) -> str:
    return re.sub(r"\.[A-Za-z0-9]{1,5}$", "", file_name.strip().replace("\\", "/").split("/")[-1]) or "file"


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


def _clean_headers(raw: list[object]) -> list[str]:
    """Trimmed, never empty, never duplicated (a second `Name` becomes `Name_2`), so every column can be addressed."""
    seen: dict[str, int] = {}
    headers: list[str] = []
    for i, value in enumerate(raw, start=1):
        text = str(value).strip() if value is not None else ""
        base = text[:150] or f"column_{i}"
        count = seen.get(base.lower(), 0) + 1
        seen[base.lower()] = count
        headers.append(base if count == 1 else f"{base}_{count}")
    return headers


# Identifiers must survive typing untouched: "000123" (an account number), "+27821234567" (a phone number)
# and 19-digit references are TEXT, never numbers, or a join to the same ids elsewhere would silently fail.
# So: no leading zeros except a lone "0", no leading "+", and at most 18 digits for a whole number.
_INT = re.compile(r"^-?(0|[1-9]\d{0,17})$")
_FLOAT = re.compile(r"^-?((0|[1-9]\d{0,17})(\.\d*)?|\.\d+)([eE][+-]?\d+)?$")
_DATE = re.compile(r"^\d{4}-\d{2}-\d{2}$")
_DATETIME = re.compile(r"^\d{4}-\d{2}-\d{2}[ T]\d{2}:\d{2}(:\d{2})?(\.\d+)?(Z|[+-]\d{2}:?\d{2})?$")
_BOOL = {"true": True, "false": False}


def infer_type(values: list[object]) -> str:
    present = [v for v in values if v is not None and v != ""]
    if not present:
        return "string"
    if all(isinstance(v, bool) for v in present) or all(isinstance(v, str) and v.strip().lower() in _BOOL for v in present):
        return "boolean"
    if all(isinstance(v, int) and not isinstance(v, bool) for v in present) or all(isinstance(v, str) and _INT.match(v.strip()) for v in present):
        return "integer"
    if all(isinstance(v, (int, float)) and not isinstance(v, bool) for v in present) or all(isinstance(v, str) and _FLOAT.match(v.strip()) for v in present):
        return "double"
    if all(isinstance(v, (date, datetime)) for v in present) or all(isinstance(v, str) and (_DATE.match(v.strip()) or _DATETIME.match(v.strip())) for v in present):
        return "date"
    return "string"


def _coerce(value: object, kind: str) -> object:
    if value is None or value == "":
        return None
    if isinstance(value, str):
        value = value.strip()
        if value == "":
            return None
    try:
        if kind == "integer":
            return int(value)
        if kind == "double":
            return float(value)
        if kind == "boolean":
            return value if isinstance(value, bool) else _BOOL[str(value).strip().lower()]
        if kind == "date":
            if isinstance(value, datetime):
                # A date-only Excel cell arrives as a datetime at midnight: keep it a plain date.
                return value.date().isoformat() if value.time() == datetime.min.time() and value.tzinfo is None else value.isoformat()
            if isinstance(value, date):
                return value.isoformat()
            return str(value)  # ISO text: the rule engine parses dates itself
    except (ValueError, KeyError, OverflowError):
        return str(value)[:MAX_CELL_CHARS]
    return str(value)[:MAX_CELL_CHARS] if not isinstance(value, (int, float, bool)) else value


def _row_limit(column_count: int) -> int:
    return max(1, min(MAX_ROWS, MAX_CELLS // max(column_count, 1)))


def _finish(name: str, headers: list[str], raw_rows: list[list[object]], truncated: bool) -> ParsedTable:
    # Pad first: zip() stops at the SHORTEST row, which would silently drop a column's values (and its type).
    width = len(headers)
    raw_rows = [list(r) + [None] * (width - len(r)) for r in raw_rows]
    columns = list(zip(*raw_rows)) if raw_rows else [() for _ in headers]
    types = {h: infer_type(list(columns[i]) if i < len(columns) else []) for i, h in enumerate(headers)}
    rows = []
    for raw in raw_rows:
        rows.append({h: _coerce(raw[i] if i < len(raw) else None, types[h]) for i, h in enumerate(headers)})
    return ParsedTable(name=name, headers=headers, rows=rows, types=types, truncated=truncated)


def table_from_records(name: str, records: list[dict], *, truncated: bool = False) -> ParsedTable:
    """A typed table from already-parsed records (an API's rows). Columns are the keys, in first-seen order;
    each column is typed and coerced exactly as a file's columns are, so a rule compares numbers with numbers
    whatever the source, and a value an API sent as text but which is plainly a number is treated as one."""
    headers: list[str] = []
    seen: set[str] = set()
    for record in records:
        for key in record:
            if key not in seen:
                seen.add(key)
                headers.append(key)
    if len(headers) > MAX_COLUMNS:
        headers = headers[:MAX_COLUMNS]
    limit = _row_limit(len(headers))
    if len(records) > limit:
        records, truncated = records[:limit], True
    return _finish(name, headers, [[record.get(h) for h in headers] for record in records], truncated)


def parse_csv(file_name: str, data: bytes) -> list[ParsedTable]:
    text = _decode(data)
    if not text.strip():
        raise FileParseError("The file is empty.")
    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel_tab if file_name.lower().endswith(".tsv") else csv.excel
    reader = csv.reader(io.StringIO(text, newline=""), dialect)
    raw_rows: list[list[object]] = []
    headers: list[str] | None = None
    truncated = False
    try:
        for record in reader:
            if not any((c or "").strip() for c in record):
                continue  # blank line
            if headers is None:
                if len(record) > MAX_COLUMNS:
                    raise FileParseError(f"The file has more than {MAX_COLUMNS} columns.")
                headers = _clean_headers(record)
                continue
            if len(raw_rows) >= _row_limit(len(headers)):
                truncated = True
                break
            raw_rows.append([c[:MAX_CELL_CHARS] for c in record[: len(headers)]])
    except csv.Error as exc:
        raise FileParseError("The file is not valid CSV.") from exc
    if headers is None:
        raise FileParseError("The file has no header row.")
    return [_finish(table_stem(file_name), headers, raw_rows, truncated)]


def _check_zip(data: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            total = sum(info.file_size for info in archive.infolist())
            if total > MAX_UNCOMPRESSED_XLSX_BYTES:
                raise FileParseError("The workbook expands to an unreasonable size and was refused.")
    except zipfile.BadZipFile as exc:
        raise FileParseError("The file is not a valid .xlsx workbook.") from exc


def parse_excel(file_name: str, data: bytes) -> list[ParsedTable]:
    _check_zip(data)
    from openpyxl import load_workbook  # local import: only Excel uploads pay for it

    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 — openpyxl raises many types for damaged files
        raise FileParseError("The workbook could not be read.") from exc
    stem = table_stem(file_name)
    tables: list[ParsedTable] = []
    try:
        for sheet in workbook.worksheets:
            headers: list[str] | None = None
            raw_rows: list[list[object]] = []
            truncated = False
            for record in sheet.iter_rows(values_only=True):
                if not any(c is not None and str(c).strip() != "" for c in record):
                    continue
                if headers is None:
                    trimmed = list(record)
                    while trimmed and (trimmed[-1] is None or str(trimmed[-1]).strip() == ""):
                        trimmed.pop()
                    if len(trimmed) > MAX_COLUMNS:
                        raise FileParseError(f"Sheet '{sheet.title}' has more than {MAX_COLUMNS} columns.")
                    headers = _clean_headers(trimmed)
                    continue
                if len(raw_rows) >= _row_limit(len(headers)):
                    truncated = True
                    break
                raw_rows.append([c if not isinstance(c, str) else c[:MAX_CELL_CHARS] for c in record[: len(headers)]])
            if headers:
                tables.append(_finish(f"{stem} / {sheet.title}" if len(workbook.worksheets) > 1 else stem, headers, raw_rows, truncated))
    finally:
        workbook.close()
    if not tables:
        raise FileParseError("The workbook has no sheet with a header row.")
    return tables


def parse_file(file_name: str, data: bytes) -> list[ParsedTable]:
    if len(data) > MAX_FILE_BYTES:
        raise FileParseError(f"The file is larger than the {MAX_FILE_BYTES // (1024 * 1024)} MB limit.")
    if not data:
        raise FileParseError("The file is empty.")
    return parse_excel(file_name, data) if file_kind(file_name) == "excel" else parse_csv(file_name, data)
